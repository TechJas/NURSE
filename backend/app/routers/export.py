# backend/app/routers/export.py
# PDF and DOCX export for nurse duty schedules

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO
from datetime import date as _date

from app.database.session import SessionLocal
from app.database.models import Nurse, ShiftAssignment
from app.utils.validators import get_current_nurse_user

router = APIRouter(prefix="/nurses/export", tags=["Export"])

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@router.get("/tasks")
def export_tasks(
    format: str = "pdf",
    db: Session = Depends(get_db),
    current_nurse: dict = Depends(get_current_nurse_user)
):
    """Export the nurse's assigned tasks as PDF or DOCX."""
    import uuid as _uuid
    nurse_id = _uuid.UUID(current_nurse["id"])
    nurse = db.query(Nurse).filter(Nurse.id == nurse_id).first()
    if not nurse:
        raise HTTPException(404, "Nurse not found")

    tasks = db.query(ShiftAssignment).filter(
        ShiftAssignment.nurse_id == nurse_id,
        ShiftAssignment.task_title != None,
    ).order_by(ShiftAssignment.shift_date.desc()).all()

    if format.lower() == "docx":
        buf = _generate_docx(nurse, tasks)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename=duty_schedule_{nurse.full_name.replace(' ','_')}.docx"}
        )
    else:
        buf = _generate_pdf(nurse, tasks)
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=duty_schedule_{nurse.full_name.replace(' ','_')}.pdf"}
        )


def _generate_pdf(nurse, tasks):
    """Generate a professional PDF duty schedule using ReportLab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm, topMargin=20*mm, bottomMargin=20*mm)
    styles = getSampleStyleSheet()
    elements = []

    # Title style
    title_style = ParagraphStyle('TitleStyle', parent=styles['Title'],
        fontSize=18, textColor=colors.HexColor('#0f2f6e'), spaceAfter=6)
    subtitle_style = ParagraphStyle('SubStyle', parent=styles['Normal'],
        fontSize=11, textColor=colors.HexColor('#4b5563'), spaceAfter=16)
    footer_style = ParagraphStyle('FooterStyle', parent=styles['Normal'],
        fontSize=9, textColor=colors.HexColor('#9ca3af'))

    # Header
    elements.append(Paragraph("🏥 Carenium Health", title_style))
    elements.append(Paragraph(
        f"<b>Duty Schedule</b> — {nurse.full_name} | {nurse.specialization or 'General'} | "
        f"CARE ID: {nurse.care_id or 'N/A'}", subtitle_style))
    elements.append(Spacer(1, 6))

    # Task table
    header = ['#', 'Date', 'Task', 'Type', 'Priority', 'Ward', 'Time', 'Status']
    data = [header]

    for i, t in enumerate(tasks, 1):
        data.append([
            str(i),
            str(t.shift_date),
            t.task_title or '—',
            t.task_type or '—',
            (t.priority or 'routine').capitalize(),
            t.ward or t.location or '—',
            t.shift_time or '—',
            (t.task_status or 'assigned').capitalize(),
        ])

    if len(data) == 1:
        data.append(['—', '—', 'No tasks assigned', '—', '—', '—', '—', '—'])

    col_widths = [20, 65, 100, 75, 55, 70, 45, 65]
    table = Table(data, colWidths=col_widths, repeatRows=1)

    # Priority colour mapping
    PRIO_COLORS = {
        'Emergency': colors.HexColor('#fef2f2'),
        'Urgent':    colors.HexColor('#fffbeb'),
        'Routine':   colors.HexColor('#f0fdf4'),
    }

    style_commands = [
        ('BACKGROUND',   (0, 0), (-1, 0), colors.HexColor('#0f2f6e')),
        ('TEXTCOLOR',    (0, 0), (-1, 0), colors.white),
        ('FONTNAME',     (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE',     (0, 0), (-1, 0), 9),
        ('FONTSIZE',     (0, 1), (-1, -1), 8),
        ('ALIGN',        (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID',         (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ('TOPPADDING',   (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING',(0, 0), (-1, -1), 5),
        ('LEFTPADDING',  (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]

    # Colour-code priority column (col 4)
    for row_idx in range(1, len(data)):
        prio = data[row_idx][4]
        bg = PRIO_COLORS.get(prio)
        if bg:
            style_commands.append(('BACKGROUND', (4, row_idx), (4, row_idx), bg))

    table.setStyle(TableStyle(style_commands))
    elements.append(table)
    elements.append(Spacer(1, 20))

    # Footer
    elements.append(Paragraph(
        f"Generated on {_date.today().strftime('%d %B %Y')} · "
        f"Total tasks: {len(tasks)} · Carenium Health Platform",
        footer_style))

    doc.build(elements)
    buf.seek(0)
    return buf


def _generate_docx(nurse, tasks):
    """Generate a professional DOCX duty schedule using python-docx."""
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Title
    title = doc.add_heading('Carenium Health', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x2F, 0x6E)

    # Subtitle
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f'Duty Schedule — {nurse.full_name}')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x37, 0x41, 0x5F)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run(f'Specialization: {nurse.specialization or "General"} | CARE ID: {nurse.care_id or "N/A"}')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph()  # spacer

    # Table
    headers = ['#', 'Date', 'Task', 'Type', 'Priority', 'Ward', 'Time', 'Status']
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    # Data rows
    for idx, t in enumerate(tasks, 1):
        row = table.add_row()
        vals = [
            str(idx),
            str(t.shift_date),
            t.task_title or '—',
            t.task_type or '—',
            (t.priority or 'routine').capitalize(),
            t.ward or t.location or '—',
            t.shift_time or '—',
            (t.task_status or 'assigned').capitalize(),
        ]
        for i, v in enumerate(vals):
            cell = row.cells[i]
            cell.text = v
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)

    if not tasks:
        row = table.add_row()
        row.cells[0].text = '—'
        row.cells[2].text = 'No tasks assigned'

    doc.add_paragraph()

    # Footer
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        f'Generated on {_date.today().strftime("%d %B %Y")} · '
        f'Total tasks: {len(tasks)} · Carenium Health Platform'
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
